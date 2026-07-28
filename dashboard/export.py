"""Excel export -- one workbook per period, built from the same KPI dicts
the dashboard's own cards render (get_summary, get_sla_kpis, etc.) plus a raw
ticket-level sheet, so "download the data" means the same numbers, not a
second definition that can drift from what's on screen.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from core.orm import NpsResponse, Ticket
from dashboard import backline, customers, frontline, utils
from dashboard.customer_matching import match_names

# Mirrors frontend/src/format.ts's hubspotTicketUrl -- same portal, same
# ticket object type ID (0-5), verified against this portal's own API
# response. Kept in sync manually; the two are far enough apart (Python
# export vs. TS UI) that sharing a constant isn't worth the coupling.
_HUBSPOT_PORTAL_ID = 2586902


def _hubspot_ticket_url(ticket_id: str) -> str:
    return f"https://app.hubspot.com/contacts/{_HUBSPOT_PORTAL_ID}/record/0-5/{ticket_id}"

_TICKET_COLUMNS = [
    "ticket_id",
    "subject",
    "pipeline_label",
    "stage_label",
    "canonical_status",
    "module",
    "primary_category",
    "customer_name",
    "priority",
    "derived_priority",
    "priority_inferred",
    "owner_name",
    "source_type",
    "reporter_contact_name",
    "created_at",
    "closed_at",
    "sla_first_response_status",
    "sla_close_status",
    "final_resolution",
    "resolution_bucket",
    "actionable",
    "fcr",
    "backline_engineer",
    "backline_path",
    "jira_link",
    "time_to_first_agent_reply_hours",
    "time_to_close_hours",
]


def _cell_value(value):
    """openpyxl can't write tz-aware datetimes or lists directly."""
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    if hasattr(value, "isoformat"):
        return value.replace(tzinfo=None).isoformat() if value.tzinfo else value.isoformat()
    return value


def _autosize(ws: Worksheet) -> None:
    for i, column_cells in enumerate(ws.columns, start=1):
        length = max((len(str(c.value)) for c in column_cells if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(length + 2, 60)


def _write_table(ws: Worksheet, rows: list[dict]) -> None:
    if not rows:
        ws.append(["No data for this period"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([_cell_value(row.get(h)) for h in headers])
    _autosize(ws)


def _flatten(prefix: str, value, out: list[tuple[str, object]]) -> None:
    """Turns a KPI function's nested dict (breakdowns, sub-percentages) into
    flat "dotted.key -> value" rows -- one routine covers every KPI dict here
    instead of a bespoke formatter per metric."""
    if isinstance(value, dict):
        for k, v in value.items():
            _flatten(f"{prefix}.{k}" if prefix else str(k), v, out)
    else:
        out.append((prefix, _cell_value(value)))


def _write_kv_section(ws: Worksheet, title: str, data: dict) -> None:
    ws.append([title])
    rows: list[tuple[str, object]] = []
    _flatten("", data, rows)
    for key, value in rows:
        ws.append([key, value])
    ws.append([])


async def _fetch_tickets(session: AsyncSession, period: str) -> list[dict]:
    period_start, period_end = utils.resolve_period(period)
    rows = await session.execute(
        select(*[getattr(Ticket, c) for c in _TICKET_COLUMNS])
        .where(Ticket.created_at.between(period_start, period_end))
        .order_by(Ticket.created_at.desc())
    )
    return [dict(zip(_TICKET_COLUMNS, row)) for row in rows.all()]


_NPS_COLUMNS = [
    "response_id",
    "email",
    "score",
    "text",
    "created_at",
    "tags",
    "excluded_from_calculations",
]


def _nps_company(properties: dict | None) -> str:
    """Wootric's "company" end-user property -- same field
    dashboard.utils.get_nps uses for account_name (verified live
    2026-07-13, populated on 76/77 responses)."""
    return (properties or {}).get("company") or "Unknown"


async def _fetch_nps_responses(session: AsyncSession, period: str) -> list[dict]:
    period_start, period_end = utils.resolve_period(period)
    rows = await session.execute(
        select(*[getattr(NpsResponse, c) for c in _NPS_COLUMNS], NpsResponse.properties)
        .where(NpsResponse.created_at.between(period_start, period_end))
        .order_by(NpsResponse.created_at.desc())
    )
    results = []
    for *values, properties in rows.all():
        row = dict(zip(_NPS_COLUMNS, values))
        row["company"] = _nps_company(properties)
        results.append(row)
    return results


async def build_export_workbook(session: AsyncSession, period: str) -> bytes:
    tickets = await _fetch_tickets(session, period)
    agents = await utils.get_agent_kpis(session, period)

    summary = await utils.get_summary(session, period)
    sla = await utils.get_sla_kpis(session, period)
    csat = await utils.get_csat(session, period)
    data_quality = await utils.get_data_quality(session, period)
    backline_overview = await backline.get_backline_overview(session, period)
    frt = await frontline.get_frontline_frt(session, period)
    fcr = await frontline.get_frontline_fcr(session, period)

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"
    for title, data in [
        ("Summary", summary),
        ("SLA", sla),
        ("CSAT", csat),
        ("Data quality", data_quality),
        ("Backline overview", backline_overview),
        # Per-owner breakdown is covered by the Agents sheet instead.
        ("Frontline FRT", {k: v for k, v in frt.items() if k != "by_owner"}),
        ("Frontline FCR", fcr),
    ]:
        _write_kv_section(ws_summary, title, data)
    _autosize(ws_summary)

    _write_table(wb.create_sheet("Tickets"), tickets)
    _write_table(wb.create_sheet("Agents"), agents)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def build_nps_workbook(session: AsyncSession, period: str) -> bytes:
    """NPS-only counterpart to build_export_workbook -- raw response rows,
    split out so the NPS card's own download button doesn't require pulling
    the whole helpdesk export."""
    nps_responses = await _fetch_nps_responses(session, period)

    wb = Workbook()
    ws = wb.active
    ws.title = "NPS Responses"
    _write_table(ws, nps_responses)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# Styling for the "customer name list" report below -- matches the
# reference "Customer tickets raised.xlsx" workbook this feature replicates
# (title row, blank row, blue header row, frozen at row 4).
_REPORT_HEADER_FILL = PatternFill("solid", fgColor="305496")
_REPORT_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _write_report_sheet(
    ws: Worksheet,
    title: str,
    headers: list[str],
    rows: list[list],
    decimal_columns: frozenset[int] = frozenset(),
) -> None:
    ws.append([title])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws.append([])
    ws.append(headers)
    for cell in ws[3]:
        cell.font = _REPORT_HEADER_FONT
        cell.fill = _REPORT_HEADER_FILL
    if not rows:
        ws.append(["No matches"])
    for row in rows:
        ws.append(row)
        for col in decimal_columns:
            ws.cell(row=ws.max_row, column=col).number_format = "0.0"
    ws.freeze_panes = "A4"
    _autosize(ws)


def _format_hms(hours: float) -> str:
    total_seconds = round(hours * 3600)
    h, rem = divmod(total_seconds, 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


async def build_customer_list_workbook(session: AsyncSession, input_names: list[str]) -> bytes:
    """Given a free-text list of company names, matches them (best-effort,
    see dashboard/customer_matching.py) against ticket customer_name values
    and builds the same 3-sheet report as the reference "Customer tickets
    raised.xlsx": per-company ticket-count/resolution summary, raw matched
    tickets, and companies with no match. Not period-scoped -- this is an
    ad-hoc "how does this account list look across all tickets" pull, not
    one of the dashboard's period-based KPIs."""
    ticket_names = (await customers.get_customer_names(session))["customer_names"]
    matched, unmatched = match_names(input_names, ticket_names)

    by_blackops: dict[str, list[tuple[str, float | None]]] = {}
    if matched:
        result = await session.execute(
            select(Ticket.customer_name, Ticket.ticket_id, Ticket.time_to_close_hours).where(
                Ticket.customer_name.in_(matched.keys())
            )
        )
        for blackops_name, ticket_id, hours in result.all():
            by_blackops.setdefault(blackops_name, []).append((ticket_id, hours))

    summary_rows = []
    for blackops_name, aliases in matched.items():
        tickets = by_blackops.get(blackops_name, [])
        hours_list = [h for _, h in tickets if h is not None]
        summary_rows.append(
            [
                ", ".join(aliases),
                blackops_name,
                len(tickets),
                round(sum(hours_list) / len(hours_list), 1) if hours_list else None,
                round(min(hours_list), 1) if hours_list else None,
                round(max(hours_list), 1) if hours_list else None,
            ]
        )
    summary_rows.sort(key=lambda r: r[2], reverse=True)

    detail_rows = [
        [blackops_name, ticket_id, _format_hms(hours), round(hours, 1)]
        for blackops_name, tickets in sorted(by_blackops.items())
        for ticket_id, hours in tickets
        if hours is not None
    ]

    no_ticket_rows = [[name, ""] for name in unmatched]

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = _safe_sheet_title("Company Summary")
    _write_report_sheet(
        ws_summary,
        "Companies With Tickets — sorted by ticket count",
        ["Company Name(s)", "Blackops Name", "Ticket Count", "Avg Resolution (hrs)", "Min Resolution (hrs)", "Max Resolution (hrs)"],
        summary_rows,
        decimal_columns=frozenset({4, 5, 6}),
    )
    _write_report_sheet(
        wb.create_sheet(_safe_sheet_title("Copy of Matched Tickets (Raw)")),
        "All Matched Tickets — Detail",
        ["Blackops Name", "Ticket ID", "Resolution time in Hours", "Resolution Hours (decimal)"],
        detail_rows,
        decimal_columns=frozenset({4}),
    )
    _write_report_sheet(
        wb.create_sheet(_safe_sheet_title("Copy of No Tickets Found")),
        "Target-List Companies With No Matching Tickets",
        ["Company Name(s)", "Blackops Name"],
        no_ticket_rows,
    )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# (ORM attribute, column header) pairs, in the order the "tickets by
# customer" export should show them -- ticket_id is handled separately since
# it needs a hyperlink, not a plain value.
_CUSTOMER_TICKET_DETAIL_COLUMNS = [
    ("owner_name", "Ticket Owner"),
    ("backline_engineer", "Backline Engineer"),
    ("created_at", "Created At"),
    ("closed_at", "Closed At"),
    ("canonical_status", "Ticket Status"),
    ("time_to_first_agent_reply_hours", "FRT (hrs)"),
    ("fcr", "FCR"),
    ("time_to_close_hours", "Resolution Time (hrs)"),
    ("final_resolution", "Final Resolution"),
    ("module", "Module"),
    ("primary_category", "Category"),
    ("sub_category", "Sub-Category"),
]


def _safe_sheet_title(name: str) -> str:
    """Excel sheet titles: max 31 chars, no : \\ / ? * [ ]."""
    for ch in ':\\/?*[]':
        name = name.replace(ch, " ")
    return name[:31]


async def build_customer_ticket_detail_workbook(
    session: AsyncSession, customer_name: str, period: str
) -> bytes:
    """One sheet: every ticket for a single identified customer/account in
    the period, ticket-level (not aggregated) -- for support/account leads
    who need the actual ticket list, not just the counts
    build_customer_counts_workbook gives across all customers."""
    period_start, period_end = utils.resolve_period(period)

    columns = ["ticket_id"] + [attr for attr, _ in _CUSTOMER_TICKET_DETAIL_COLUMNS]
    rows = await session.execute(
        select(*[getattr(Ticket, c) for c in columns])
        .where(Ticket.customer_name == customer_name, Ticket.created_at.between(period_start, period_end))
        .order_by(Ticket.created_at.desc())
    )
    tickets = rows.all()

    wb = Workbook()
    ws = wb.active
    ws.title = _safe_sheet_title(customer_name)

    headers = ["HubSpot Ticket ID"] + [label for _, label in _CUSTOMER_TICKET_DETAIL_COLUMNS]
    heading = f"{customer_name} tickets: {period_start.date().isoformat()} – {period_end.date().isoformat()}"
    ws.append([heading])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.append(headers)

    if not tickets:
        ws.append(["No tickets for this customer in this period"])
    for row in tickets:
        ticket_id = row[0]
        ws.append([ticket_id] + [_cell_value(v) for v in row[1:]])
        id_cell = ws.cell(row=ws.max_row, column=1)
        id_cell.hyperlink = _hubspot_ticket_url(ticket_id)
        id_cell.style = "Hyperlink"

    _autosize(ws)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def build_customer_counts_workbook(session: AsyncSession, period: str) -> bytes:
    """Single-sheet workbook: issues reported per customer for the period --
    the same aggregate /customers/volume shows, but the full list instead of
    top-N, plus the no-account bucket and a total."""
    period_start, period_end = utils.resolve_period(period)
    in_period = Ticket.created_at.between(period_start, period_end)

    rows = await session.execute(
        select(Ticket.customer_name, func.count())
        .where(in_period, Ticket.customer_name.isnot(None))
        .group_by(Ticket.customer_name)
        .order_by(func.count().desc())
    )
    identified = rows.all()
    no_account_count = await session.scalar(
        select(func.count()).where(in_period, Ticket.customer_name.is_(None))
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Customer Issues"
    ws.append(["Customer", "Issues reported"])
    for name, count in identified:
        ws.append([name, count])
    ws.append(["(No account)", no_account_count or 0])
    ws.append(["Total", sum(count for _, count in identified) + (no_account_count or 0)])
    _autosize(ws)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

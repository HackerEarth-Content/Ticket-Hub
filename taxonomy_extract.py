"""Standalone export: module -> category -> sub-category taxonomy, as an
Excel file with collapsible outline grouping.

Module->category comes from hubspot_pipeline/category_taxonomy.json
(resolve_module), same as the dashboard's module distribution chart.
Category->sub-category isn't config'd anywhere -- HubSpot's sub_category
field is independent of the multi-select hs_ticket_category field -- so
which sub-categories belong to which category is derived empirically from
which combinations actually appear together on a ticket.

Run: python taxonomy_extract.py [output.xlsx]
"""

from __future__ import annotations

import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from core.config import settings
from core.orm import Ticket
from hubspot_pipeline.category_map import UNCATEGORIZED_MODULE, resolve_module

_UNCATEGORIZED_SUB = "Uncategorized"

_MODULE_FILL = PatternFill("solid", fgColor="1F2937")
_CATEGORY_FILL = PatternFill("solid", fgColor="E5E7EB")
_MODULE_FONT = Font(bold=True, color="FFFFFF")
_CATEGORY_FONT = Font(bold=True)
_HEADER_FONT = Font(bold=True)


def build_taxonomy(rows: list[tuple[list[str], str | None]]) -> dict[str, dict[str, set[str]]]:
    """rows: (categories, sub_category) per ticket -> module -> category -> sub-categories seen."""
    tree: dict[str, dict[str, set[str]]] = defaultdict(lambda: defaultdict(set))
    for categories, sub_category in rows:
        sub = sub_category or _UNCATEGORIZED_SUB
        if not categories:
            tree[UNCATEGORIZED_MODULE][UNCATEGORIZED_MODULE].add(sub)
            continue
        for category in categories:
            tree[resolve_module([category])][category].add(sub)
    return tree


def _fetch_rows() -> list[tuple[list[str], str | None]]:
    url = settings.DATABASE_URL.replace("postgresql://", "postgresql+psycopg://")
    engine = create_engine(url)
    with Session(engine) as session:
        return session.execute(select(Ticket.categories, Ticket.sub_category)).all()


def write_excel(tree: dict[str, dict[str, set[str]]], out_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Taxonomy"
    ws.sheet_properties.outlinePr.summaryBelow = False

    ws.append(["Module", "Category", "Sub-category"])
    for cell in ws[1]:
        cell.font = _HEADER_FONT
    ws.freeze_panes = "A2"

    for module, categories in sorted(tree.items()):
        row = ws.max_row + 1
        ws.append([module, "", ""])
        for cell in ws[row]:
            cell.fill = _MODULE_FILL
            cell.font = _MODULE_FONT

        for category, sub_categories in sorted(categories.items()):
            row = ws.max_row + 1
            ws.append(["", category, ""])
            for cell in ws[row]:
                cell.fill = _CATEGORY_FILL
                cell.font = _CATEGORY_FONT
            ws.row_dimensions[row].outlineLevel = 1

            for sub_category in sorted(sub_categories):
                row = ws.max_row + 1
                ws.append(["", "", sub_category])
                ws.row_dimensions[row].outlineLevel = 2

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 28

    wb.save(out_path)


def main() -> None:
    out_path = sys.argv[1] if len(sys.argv) > 1 else "taxonomy_extract.xlsx"
    tree = build_taxonomy(_fetch_rows())
    write_excel(tree, out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
